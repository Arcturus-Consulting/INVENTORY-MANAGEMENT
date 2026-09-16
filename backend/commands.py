"""Deterministic command engine for the ORION validator UI."""

from __future__ import annotations

import difflib
import logging
import re
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

logger = logging.getLogger("orion_validator.commands")

COMMANDS = {
    "HELP": "Show commands and the supported workflow.",
    "STATUS": "Show Oracle readiness, the loaded report, validation, and sign-off state.",
    "VALIDATE": "Run local arithmetic and live Oracle checks.",
    "SUMMARY": "Show the latest scorecard and live-check coverage.",
    "EXCEPTIONS": "List rows with failed checks.",
    "DETAILS": "Show one item/valuation-unit result. Usage: DETAILS <item> [@ <valuation unit>]",
    "SIGNOFF": "Show sign-off status or approve a fully passed run. Usage: SIGNOFF <approver>",
    "EXPORT": "Generate an auditable Excel validation workbook.",
    "RESET": "Clear the current report, validation, export, and sign-off.",
}

COMMAND_NAMES = list(COMMANDS)
STATUS_BADGES = {
    "pass": "🟢 PASS",
    "fail": "🔴 FAIL",
    "incomplete": "🟠 INCOMPLETE",
    "skipped": "🟠 SKIPPED",
}
CHECK_LABELS = {
    "rollforward": "① Quantity roll-forward",
    "value_calc": "② Inventory value calculation",
    "oracle_qty": "③ Oracle live on-hand cross-check",
    "oracle_cost": "④ Oracle live cost cross-check",
}


def _help_text() -> str:
    lines = ["**Available commands:**"]
    for name, description in COMMANDS.items():
        lines.append(f"- `{name}` — {description}")
    lines.append(
        "\nAttach an Oracle Period Inventory Valuation Report (.xlsx, .csv, or text-based .pdf), "
        "then type `VALIDATE`."
    )
    return "\n".join(lines)


def _status_text(session: dict) -> str:
    from oracle_client import get_oracle_connection_status

    oracle = get_oracle_connection_status()
    lines = [
        f"**Oracle live checks:** {'🟢 Configured' if oracle['configured'] else '🟠 Not configured'} — "
        f"{oracle['message']}"
    ]
    report = session.get("report")
    if report:
        source = session.get("source_filename", "uploaded report")
        lines.append(f"**Report loaded:** {source} — {report['row_count']} item rows.")
        for key in ("Cost Organization", "Cost Book", "Period", "Currency", "Period Start", "Period End"):
            if report.get("metadata", {}).get(key):
                lines.append(f"- {key}: {report['metadata'][key]}")
    else:
        lines.append("**Report loaded:** none. Attach a report to begin.")

    validation = session.get("validation")
    if validation:
        lines.append(
            f"**Latest validation:** {STATUS_BADGES[validation['overall_status']]} — "
            f"run `{validation['run_id']}`."
        )
    else:
        lines.append("**Latest validation:** none.")

    signoff = session.get("signoff")
    if signoff:
        lines.append(f"**Sign-off:** approved by {signoff['approver']} at {signoff['signed_at']}.")
    else:
        lines.append("**Sign-off:** not approved.")
    return "\n".join(lines)


def _summary_text(session: dict) -> dict:
    validation = session.get("validation")
    if not validation:
        return {"type": "text", "message": "No validation has run yet. Attach a report and type `VALIDATE`."}

    live = validation["live_coverage"]
    lines = [
        f"## {STATUS_BADGES[validation['overall_status']]}",
        "",
        f"- Item rows: **{validation['row_count']}**",
        f"- Fully passed rows: **{validation['pass_count']}**",
        f"- Rows with failures: **{validation['fail_count']}**",
        f"- Incomplete rows (one or more skipped checks): **{validation['incomplete_count']}**",
        f"- Live Oracle checks completed: **{live['completed']} / {live['expected']}**",
        f"- Computed Inventory Value total: **{validation['computed_value_sum']:,.2f}**",
        "",
        f"**Total tie-out:** {validation['total_check']['detail']}",
        "",
        "**Check coverage:**",
    ]
    for key, label in CHECK_LABELS.items():
        counts = validation["check_summary"][key]
        lines.append(
            f"- {label}: {counts['pass']} pass / {counts['fail']} fail / {counts['skipped']} skipped"
        )

    if validation["overall_status"] == "fail":
        lines.append("\nType `EXCEPTIONS`, then `DETAILS <item>` to investigate.")
    elif validation["overall_status"] == "incomplete":
        lines.append(
            "\nLocal checks passed, but one or more required evidence checks were skipped. "
            "Review `STATUS` and the item details before treating this run as signed off."
        )
    else:
        lines.append("\nAll checks passed. Type `SIGNOFF <approver>` and then `EXPORT`.")
    return {"type": "text", "message": "\n".join(lines)}


def _exceptions_text(session: dict) -> dict:
    validation = session.get("validation")
    if not validation:
        return {"type": "text", "message": "No validation has run yet. Type `VALIDATE` first."}
    failing = [result for result in validation["results"] if result["status"] == "fail"]
    if not failing:
        if validation["overall_status"] == "incomplete":
            return {
                "type": "text",
                "message": "No failed rows, but the run is incomplete because live checks were skipped. Type `SUMMARY` for coverage.",
            }
        return {"type": "text", "message": "🟢 No exceptions. Every required check passed."}

    lines = [f"**{len(failing)} row(s) with exceptions:**\n"]
    for result in failing:
        failed = [CHECK_LABELS[name] for name, check in result["checks"].items() if check["status"] == "fail"]
        lines.append(
            f"- `{result['item']}` @ `{result['valuation_unit']}` — {result['description']}  \n"
            f"  Failed: {', '.join(failed)}"
        )
    lines.append("\nType `DETAILS <item>` for the full breakdown.")
    return {"type": "text", "message": "\n".join(lines)}


def _parse_detail_selector(argument: str) -> tuple[str, str]:
    item, separator, valuation_unit = argument.partition("@")
    return item.strip(), valuation_unit.strip() if separator else ""


def _details_text(session: dict, argument: str) -> dict:
    validation = session.get("validation")
    if not validation:
        return {"type": "text", "message": "No validation has run yet. Type `VALIDATE` first."}
    item_arg, unit_arg = _parse_detail_selector(argument)
    if not item_arg:
        return {"type": "text", "message": "Provide an item, for example `DETAILS ECM500002`."}

    matches = [
        result
        for result in validation["results"]
        if result["item"].upper() == item_arg.upper()
        and (not unit_arg or result["valuation_unit"].upper() == unit_arg.upper())
    ]
    if len(matches) > 1:
        options = "\n".join(
            f"- `DETAILS {item_arg} @ {result['valuation_unit']}`" for result in matches
        )
        return {
            "type": "text",
            "message": f"Item `{item_arg}` occurs in multiple valuation units. Choose one:\n{options}",
        }
    if not matches:
        all_items = sorted({result["item"] for result in validation["results"]})
        suggestions = difflib.get_close_matches(item_arg.upper(), [item.upper() for item in all_items], n=3, cutoff=0.5)
        suffix = f" Did you mean: {', '.join(f'`{item}`' for item in suggestions)}?" if suggestions else ""
        return {"type": "text", "message": f"Item `{item_arg}` was not found in this validation run.{suffix}"}

    result = matches[0]
    row = result["row_data"]
    lines = [
        f"## {result['item']} — {result['description']}",
        f"**Overall: {STATUS_BADGES[result['status']]}**",
        f"Valuation Unit: {result['valuation_unit']} | UOM: {row.get('uom', '')}",
        "",
        (
            f"Opening {row.get('opening_quantity')} | Receipts {row.get('receipts')} | "
            f"Issues {row.get('issues')} | Closing {row.get('closing_quantity')} | "
            f"Unit Cost {row.get('unit_cost')} | Inventory Value {row.get('inventory_value')}"
        ),
        "",
    ]
    for key, label in CHECK_LABELS.items():
        check = result["checks"][key]
        icon = {"pass": "🟢", "fail": "🔴", "skipped": "🟠"}[check["status"]]
        lines.append(f"**{icon} {label}:** {check['detail']}")
    return {"type": "text", "message": "\n".join(lines)}


def _signoff_text(session: dict, approver: str) -> dict:
    existing = session.get("signoff")
    if not approver:
        if existing:
            return {
                "type": "text",
                "message": f"Run signed off by **{existing['approver']}** at {existing['signed_at']} UTC.",
            }
        return {"type": "text", "message": "No sign-off is recorded. Usage: `SIGNOFF <approver>`."}
    validation = session.get("validation")
    if not validation:
        return {"type": "text", "message": "Run `VALIDATE` before signing off."}
    if validation["overall_status"] != "pass":
        return {
            "type": "text",
            "message": (
                f"This run is {validation['overall_status'].upper()} and cannot be signed off. "
                "Resolve failed checks or restore the skipped live evidence, then run `VALIDATE` again."
            ),
        }
    clean_name = re.sub(r"\s+", " ", approver).strip()
    if len(clean_name) > 120:
        return {"type": "text", "message": "Approver name must be 120 characters or fewer."}
    signoff = {
        "approver": clean_name,
        "signed_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "validation_run_id": validation["run_id"],
    }
    session["signoff"] = signoff
    session.pop("_last_export", None)
    return {
        "type": "text",
        "message": f"🟢 Validation run `{validation['run_id']}` signed off by **{clean_name}**.",
    }


def _style_worksheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8E4B32")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in worksheet.columns:
        width = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 11), 64)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _build_export_workbook(session: dict) -> bytes:
    validation = session["validation"]
    report = session["report"]
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("Source file", session.get("source_filename", "")),
        ("Source format", report.get("source_format", "")),
        ("Validation run ID", validation["run_id"]),
        ("Completed at UTC", validation["completed_at"]),
        ("Overall status", validation["overall_status"].upper()),
        ("Rows", validation["row_count"]),
        ("Fully passed rows", validation["pass_count"]),
        ("Failed rows", validation["fail_count"]),
        ("Incomplete rows", validation["incomplete_count"]),
        ("Computed Inventory Value", validation["computed_value_sum"]),
        ("Total tie-out status", validation["total_check"]["status"].upper()),
        ("Total tie-out detail", validation["total_check"]["detail"]),
    ]
    signoff = session.get("signoff")
    summary_rows.extend(
        [
            ("Sign-off status", "SIGNED" if signoff else "NOT SIGNED"),
            ("Approver", signoff["approver"] if signoff else ""),
            ("Signed at UTC", signoff["signed_at"] if signoff else ""),
        ]
    )
    for key, value in report.get("metadata", {}).items():
        summary_rows.append((key, value))
    summary.append(["Field", "Value"])
    for row in summary_rows:
        summary.append(list(row))
    _style_worksheet(summary)

    results_sheet = workbook.create_sheet("Validation Results")
    headers = [
        "Row",
        "Item",
        "Description",
        "Valuation Unit",
        "Overall Status",
        "Roll-forward",
        "Roll-forward Detail",
        "Value Calculation",
        "Value Calculation Detail",
        "Oracle Quantity",
        "Oracle Quantity Detail",
        "Oracle Cost",
        "Oracle Cost Detail",
    ]
    results_sheet.append(headers)
    fills = {
        "PASS": PatternFill("solid", fgColor="E2F0D9"),
        "FAIL": PatternFill("solid", fgColor="FCE4D6"),
        "INCOMPLETE": PatternFill("solid", fgColor="FFF2CC"),
    }
    for result in validation["results"]:
        checks = result["checks"]
        results_sheet.append(
            [
                result["row_number"],
                result["item"],
                result["description"],
                result["valuation_unit"],
                result["status"].upper(),
                checks["rollforward"]["status"].upper(),
                checks["rollforward"]["detail"],
                checks["value_calc"]["status"].upper(),
                checks["value_calc"]["detail"],
                checks["oracle_qty"]["status"].upper(),
                checks["oracle_qty"]["detail"],
                checks["oracle_cost"]["status"].upper(),
                checks["oracle_cost"]["detail"],
            ]
        )
        fill = fills.get(result["status"].upper())
        if fill:
            for cell in results_sheet[results_sheet.max_row]:
                cell.fill = fill
    if results_sheet.max_row > 1:
        table = Table(displayName="ValidationResults", ref=f"A1:M{results_sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        results_sheet.add_table(table)
    _style_worksheet(results_sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _unknown_command_response(raw_command: str) -> dict:
    suggestion = difflib.get_close_matches(raw_command.upper(), COMMAND_NAMES, n=1, cutoff=0.55)
    if suggestion:
        return {
            "type": "text",
            "message": f"I don't recognize `{raw_command}`. Did you mean `{suggestion[0]}`?",
        }
    return {"type": "text", "message": f"I don't recognize `{raw_command}`. Type `HELP` for the command list."}


def handle_command(raw_text: str, session: dict) -> dict:
    text = (raw_text or "").strip()
    if not text:
        return {"type": "text", "message": "Type a command, or `HELP` for the list."}
    parts = text.split(maxsplit=1)
    raw_command = parts[0]
    argument = parts[1].strip() if len(parts) > 1 else ""
    command = raw_command.upper()
    logger.info("COMMAND: %s", command)
    if command not in COMMANDS:
        return _unknown_command_response(raw_command)
    if command == "HELP":
        return {"type": "text", "message": _help_text()}
    if command == "STATUS":
        return {"type": "text", "message": _status_text(session)}
    if command == "VALIDATE":
        if not session.get("report"):
            return {"type": "text", "message": "Attach a report before running `VALIDATE`."}
        from validator import validate_report

        session["validation"] = validate_report(session["report"])
        session.pop("signoff", None)
        session.pop("_last_export", None)
        return _summary_text(session)
    if command == "SUMMARY":
        return _summary_text(session)
    if command == "EXCEPTIONS":
        return _exceptions_text(session)
    if command == "DETAILS":
        return _details_text(session, argument)
    if command == "SIGNOFF":
        return _signoff_text(session, argument)
    if command == "EXPORT":
        if not session.get("validation"):
            return {"type": "text", "message": "Run `VALIDATE` before exporting."}
        workbook_bytes = _build_export_workbook(session)
        period = session["report"].get("metadata", {}).get("Period", "")
        safe_period = re.sub(r"[^A-Za-z0-9_-]+", "_", period).strip("_")
        filename = f"ORION_Validation_{safe_period}.xlsx" if safe_period else "ORION_Validation_Report.xlsx"
        return {
            "type": "export",
            "message": "The validation workbook is ready.",
            "file_bytes": workbook_bytes,
            "filename": filename,
        }
    if command == "RESET":
        for key in ("report", "validation", "signoff", "source_filename", "_last_export", "_last_export_name"):
            session.pop(key, None)
        return {"type": "text", "message": "Session cleared. Attach a new report to begin."}
    return _unknown_command_response(raw_command)
