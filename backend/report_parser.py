"""Parse Oracle Period Inventory Valuation reports.

The parser is deterministic and supports native Oracle XLSX/CSV exports plus
text-based Oracle Analytics Publisher PDFs. Headers are discovered by content;
the report is not tied to a fixed row number or period label.
"""

from __future__ import annotations

import csv
import re
from io import BytesIO, StringIO
from pathlib import Path

import openpyxl


HEADER_ALIASES = {
    "valuation_unit": ["valuation unit"],
    "category": ["category"],
    "item": ["item"],
    "description": ["description"],
    "uom": ["uom"],
    "cost_element": ["cost element"],
    "unit_cost": ["unit cost"],
    "opening_quantity": ["opening quantity"],
    "receipts": ["receipts"],
    "issues": ["issues"],
    "closing_quantity": ["closing quantity"],
    "inventory_value": ["inventory value"],
    "prior_inventory_value": ["prior inventory value"],
    "period_activity": ["period activity"],
    "period_costed_value": ["period costed value"],
    "period_accounted_value": ["period accounted value"],
}

REQUIRED_HEADER_FIELDS = [
    "item",
    "unit_cost",
    "opening_quantity",
    "receipts",
    "issues",
    "closing_quantity",
    "inventory_value",
]

NUMERIC_FIELDS = {
    "unit_cost",
    "opening_quantity",
    "receipts",
    "issues",
    "closing_quantity",
    "inventory_value",
    "prior_inventory_value",
    "period_activity",
    "period_costed_value",
    "period_accounted_value",
}

ZERO_DEFAULT_FIELDS = {
    "opening_quantity",
    "receipts",
    "issues",
    "prior_inventory_value",
    "period_activity",
}

METADATA_LABELS = {
    "cost organization": "Cost Organization",
    "cost book": "Cost Book",
    "period": "Period",
    "inventory organization": "Inventory Organization",
    "currency": "Currency",
    "period start": "Period Start",
    "period end": "Period End",
    "summarization level": "Summarization Level",
    "valuation structure type": "Valuation Structure Type",
}


def _norm(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def _compact(value) -> str:
    return re.sub(r"[^a-z0-9]+", "", _norm(value))


def _header_matches(cell, field: str, variant: str) -> bool:
    normalized = _norm(cell)
    if normalized == variant:
        return True
    compact = _compact(cell)
    variant_compact = _compact(variant)
    if compact == variant_compact:
        return True
    # Oracle appends the period to value columns, for example
    # "Inventory Value (05-26)" and "Prior Inventory Value (04-26)".
    return field in {"inventory_value", "prior_inventory_value"} and compact.startswith(variant_compact)


def _find_header_row(rows: list[tuple | list]) -> tuple[int, dict[str, int]]:
    best_row_idx = None
    best_map: dict[str, int] = {}
    best_score = 0

    for r_idx, row in enumerate(rows):
        col_map: dict[str, int] = {}
        for c_idx, cell in enumerate(row):
            if not _norm(cell):
                continue
            for field, variants in HEADER_ALIASES.items():
                if field in col_map:
                    continue
                if any(_header_matches(cell, field, variant) for variant in variants):
                    col_map[field] = c_idx
        score = sum(field in col_map for field in REQUIRED_HEADER_FIELDS)
        if score > best_score:
            best_row_idx, best_map, best_score = r_idx, col_map, score

    if best_row_idx is None or best_score < len(REQUIRED_HEADER_FIELDS):
        raise ValueError(
            "Could not locate a valid Period Inventory Valuation Report header row. "
            "Expected Item, Unit Cost, Opening Quantity, Receipts, Issues, Closing "
            "Quantity, and Inventory Value columns."
        )
    return best_row_idx, best_map


def _to_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"\s+", "", str(value)).replace(",", "")
    if text in {"", "-"}:
        return None
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return float(text)
    except ValueError:
        return None


def _clean_text(value, *, identifier: bool = False) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if identifier:
        return re.sub(r"\s+", "", text).replace("\\", "")
    return re.sub(r"\s+", " ", text)


def _extract_metadata(rows: list[tuple | list], header_row_idx: int | None = None) -> dict:
    metadata: dict[str, str] = {}
    source_rows = rows if header_row_idx is None else rows[:header_row_idx]
    for row in source_rows:
        cells = [cell for cell in row if _norm(cell)]
        if len(cells) < 2:
            continue
        label = _norm(cells[0])
        if label in METADATA_LABELS:
            metadata[METADATA_LABELS[label]] = _clean_text(cells[1])
    return metadata


def _parse_data_rows(
    rows: list[tuple | list],
    header_row_idx: int,
    col_map: dict[str, int],
    *,
    pdf_mode: bool = False,
    existing_rows: list[dict] | None = None,
) -> tuple[list[dict], dict | None]:
    data_rows = existing_rows if existing_rows is not None else []
    total_row = None

    for row in rows[header_row_idx + 1 :]:
        if not any(_norm(cell) for cell in row):
            continue

        def get(field):
            idx = col_map.get(field)
            return row[idx] if idx is not None and idx < len(row) else None

        item = _clean_text(get("item"), identifier=pdf_mode)
        has_total_label = any(_norm(cell) in {"total", "grand total"} for cell in row)
        if not item and has_total_label:
            total_row = {
                field: _to_number(get(field)) if field in NUMERIC_FIELDS else _clean_text(get(field))
                for field in HEADER_ALIASES
            }
            continue

        if not item:
            # Oracle PDFs can continue a wrapped description on the next page as
            # a row containing only description text. Preserve it on the prior row.
            continuation = _clean_text(get("description"))
            if pdf_mode and continuation and data_rows:
                data_rows[-1]["description"] = " ".join(
                    part for part in (data_rows[-1].get("description", ""), continuation) if part
                )
            continue

        parsed_row = {}
        for field in HEADER_ALIASES:
            raw = get(field)
            if field in NUMERIC_FIELDS:
                parsed_row[field] = _to_number(raw)
            else:
                parsed_row[field] = _clean_text(
                    raw,
                    identifier=pdf_mode and field in {"valuation_unit", "item", "uom"},
                )
        for field in ZERO_DEFAULT_FIELDS:
            if parsed_row.get(field) is None:
                parsed_row[field] = 0.0
        data_rows.append(parsed_row)

    return data_rows, total_row


def _build_result(metadata: dict, rows: list[dict], total_row: dict | None, source_format: str) -> dict:
    if not rows:
        raise ValueError("No item rows were found below the report header.")
    return {
        "metadata": metadata,
        "rows": rows,
        "total_row": total_row,
        "row_count": len(rows),
        "source_format": source_format,
    }


def _parse_xlsx(file_bytes: bytes) -> dict:
    try:
        workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("The file could not be opened as an XLSX workbook.") from exc
    worksheet = workbook[workbook.sheetnames[0]]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The uploaded file is empty.")
    header_idx, col_map = _find_header_row(rows)
    parsed_rows, total_row = _parse_data_rows(rows, header_idx, col_map)
    return _build_result(_extract_metadata(rows, header_idx), parsed_rows, total_row, "xlsx")


def _decode_csv(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("The CSV encoding is not supported. Use UTF-8 or UTF-16.")


def _parse_csv(file_bytes: bytes) -> dict:
    rows = list(csv.reader(StringIO(_decode_csv(file_bytes))))
    if not rows:
        raise ValueError("The uploaded file is empty.")
    header_idx, col_map = _find_header_row(rows)
    parsed_rows, total_row = _parse_data_rows(rows, header_idx, col_map)
    return _build_result(_extract_metadata(rows, header_idx), parsed_rows, total_row, "csv")


def _parse_pdf(file_bytes: bytes) -> dict:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ValueError("PDF support is unavailable because pdfplumber is not installed.") from exc

    metadata: dict[str, str] = {}
    parsed_rows: list[dict] = []
    total_row = None
    try:
        with pdfplumber.open(BytesIO(file_bytes)) as document:
            for page in document.pages:
                for table in page.extract_tables() or []:
                    if not table:
                        continue
                    metadata.update(_extract_metadata(table))
                    try:
                        header_idx, col_map = _find_header_row(table)
                    except ValueError:
                        continue
                    parsed_rows, page_total = _parse_data_rows(
                        table,
                        header_idx,
                        col_map,
                        pdf_mode=True,
                        existing_rows=parsed_rows,
                    )
                    if page_total is not None:
                        total_row = page_total
    except Exception as exc:  # noqa: BLE001
        raise ValueError(
            "The PDF could not be parsed. Upload the text-based Oracle Analytics Publisher "
            "PDF, or use the XLSX/CSV export for scanned reports."
        ) from exc
    return _build_result(metadata, parsed_rows, total_row, "pdf")


def parse_valuation_report(file_bytes: bytes, filename: str = "report.xlsx") -> dict:
    """Parse an Oracle valuation report based on its filename extension."""
    extension = Path(filename or "").suffix.lower()
    if extension == ".xlsx":
        return _parse_xlsx(file_bytes)
    if extension == ".csv":
        return _parse_csv(file_bytes)
    if extension == ".pdf":
        return _parse_pdf(file_bytes)
    raise ValueError("Unsupported file type. Upload an .xlsx, .csv, or text-based .pdf report.")
